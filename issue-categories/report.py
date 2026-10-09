"""Excel report: top 5 problem categories per model per month, for one year, on one tab.

Each month is a block of columns; the blocks sit side by side, January on the left.

Usage: python report.py [year]      (default 2026)
Needs data/order_labels.parquet (label.py export). Writes data/reports/top-problems-<year>.xlsx.

Orders counted: OrderDate in the year, PendingApproval = 'Approved', and at least one
ShippingOrders row with ShippingStatus = 'Shipped'. Model = first 6 characters of SerialNumber.
Item cost = sum of OrderItems.Quantity * Products.Cost (UnitPrice is 0 on warranty orders).
"""

import csv
import io
import json
import sys
import urllib.request
from collections import Counter, defaultdict
from pathlib import Path

import duckdb
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from pull import SPIRIT_WEB_DB, load_env

HERE = Path(__file__).parent
DATA = HERE / "data"

QUERY = """
SELECT o.Id, CONVERT(char(7), o.OrderDate, 126) AS Month,
       CASE WHEN LEN(LTRIM(ISNULL(o.SerialNumber, ''))) >= 6 THEN LEFT(LTRIM(o.SerialNumber), 6)
            ELSE '(no serial)' END AS Model,
       ISNULL((SELECT SUM(i.Quantity * ISNULL(p.Cost, 0))
                 FROM OrderItems i LEFT JOIN Products p ON p.ProductId = i.ProductId
                WHERE i.OrderId = o.Id), 0) AS ItemCost,
       (SELECT COUNT(*) FROM OrderItems i LEFT JOIN Products p ON p.ProductId = i.ProductId
         WHERE i.OrderId = o.Id AND p.Cost IS NULL) AS ItemsWithoutCost
  FROM Orders o
 WHERE o.OrderDate >= '{year}-01-01' AND o.OrderDate < '{next_year}-01-01'
   AND o.PendingApproval = 'Approved'
   AND EXISTS (SELECT 1 FROM ShippingOrders s WHERE s.OrderId = o.Id AND s.ShippingStatus = 'Shipped')
"""

FONT = "Arial"
HEADER_FILL = PatternFill("solid", start_color="1F3864")
MODEL_FILL = PatternFill("solid", start_color="D9E1F2")
THIN = Side(style="thin", color="BFBFBF")


def fetch_orders(year):
    env = load_env()
    body = {
        "query": {"database": SPIRIT_WEB_DB, "type": "native",
                  "native": {"query": QUERY.format(year=year, next_year=year + 1)}},
        "format_rows": False,
    }
    req = urllib.request.Request(
        env["METABASE_URL"].rstrip("/") + "/api/dataset/csv",
        data=json.dumps(body).encode(),
        headers={"x-api-key": env["METABASE_API_KEY"], "Content-Type": "application/json"},
    )
    with urllib.request.urlopen(req, timeout=600) as resp:
        return list(csv.DictReader(io.StringIO(resp.read().decode())))


def categories_by_order():
    rows = duckdb.sql(f"SELECT Id, Categories FROM '{DATA / 'order_labels.parquet'}'").fetchall()
    return {str(i): c for i, c in rows}


def font(**kw):
    return Font(name=FONT, **kw)


def write_header(ws, row, headers, widths, first_col=1):
    for col, (text, width) in enumerate(zip(headers, widths), first_col):
        cell = ws.cell(row, col, text)
        cell.font = font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[row].height = 30


def write_row(ws, row, values, formats, first_col=1, bold=False, fill=None, top_border=False):
    for col, (value, fmt) in enumerate(zip(values, formats), first_col):
        cell = ws.cell(row, col, value)
        cell.font = font(bold=bold)
        if fmt:
            cell.number_format = fmt
        if fill:
            cell.fill = fill
        if top_border:
            cell.border = Border(top=THIN)


def main():
    year = int(sys.argv[1]) if len(sys.argv) > 1 else 2026
    orders = fetch_orders(year)
    labels = categories_by_order()

    count = defaultdict(Counter)
    cost = defaultdict(Counter)
    model_orders = defaultdict(Counter)
    month_orders = Counter()
    month_cost = Counter()
    for o in orders:
        month, model, item_cost = o["Month"], o["Model"], float(o["ItemCost"] or 0)
        model_orders[month][model] += 1
        month_orders[month] += 1
        month_cost[month] += item_cost
        for cat in labels.get(o["Id"]) or ["No issue text"]:
            count[(month, model)][cat] += 1
            cost[(month, model)][cat] += item_cost

    wb = Workbook()
    ws = wb.active
    ws.title = f"Top Problems {year}"
    ws["A1"] = f"Top 5 problem categories per model — {year}"
    ws["A1"].font = font(bold=True, size=14)
    ws["A2"] = ("Approved orders with a shipped shipping order. Model = first 6 digits of the serial number. "
                "An order can have up to 3 categories. Item cost = quantity × Products.Cost.")
    ws["A2"].font = font(italic=True, size=9, color="595959")

    headers = ["Model", "Rank", "Problem Category", "Orders With Category", "% of Model's Orders", "Item Cost ($)"]
    formats = [None, None, None, "#,##0", "0.0%", "$#,##0.00"]
    widths = [10, 6, 34, 11, 11, 12]
    block = len(headers) + 1
    longest = 0
    for i, month in enumerate(sorted(model_orders)):
        first = 1 + i * block
        cell = ws.cell(4, first, f"{month}    {month_orders[month]:,} orders    ${month_cost[month]:,.0f} item cost")
        cell.font = font(bold=True, size=12, color="1F3864")
        write_header(ws, 5, headers, widths, first)
        ws.column_dimensions[get_column_letter(first + len(headers))].width = 3
        row = 6
        for model, n_model in sorted(model_orders[month].items(), key=lambda kv: (-kv[1], kv[0])):
            ranked = sorted(count[(month, model)].items(),
                            key=lambda kv: (-kv[1], -cost[(month, model)][kv[0]], kv[0]))[:5]
            for rank, (cat, n) in enumerate(ranked, 1):
                write_row(ws, row, [model, rank, cat, n, n / n_model, cost[(month, model)][cat]], formats, first,
                          bold=(rank == 1), fill=MODEL_FILL if rank == 1 else None, top_border=(rank == 1))
                row += 1
        longest = max(longest, row - 1)
    ws.freeze_panes = "A6"

    out = DATA / "reports" / f"top-problems-{year}.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out}: {len(orders)} orders, {len(model_orders)} months, {longest} rows")


if __name__ == "__main__":
    main()
